"""Build the complete analytical report as a formatted Word document."""

from __future__ import annotations

import argparse
from pathlib import Path
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "report" / "FULL_REPORT.md"
DEFAULT_OUTPUT = ROOT / "report" / "Odos_Ermou_Recommendation_Report.docx"

IMAGE_PATTERN = re.compile(r"^!\[([^]]*)\]\(([^)]+)\)$")
HEADING_PATTERN = re.compile(r"^(#{1,6})\s+(.+)$")
BULLET_PATTERN = re.compile(r"^(\s*)[-*]\s+(.+)$")
NUMBER_PATTERN = re.compile(r"^(\s*)\d+\.\s+(.+)$")
TABLE_SEPARATOR = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?$")
INLINE_PATTERN = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*|\[[^]]+\]\([^)]+\)|https?://\S+)"
)


def set_cell_shading(cell, fill):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def set_repeat_table_header(row):
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(table_header)


def prevent_row_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80):
    margins = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell._tc.get_or_add_tcPr().append(margins)
    for name, value in (("top", top), ("start", start),
                        ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def add_toc(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate))
    paragraph.add_run(
        "Open in Word and update this field to display page numbers."
    )
    paragraph.add_run()._r.append(end)


def add_hyperlink(paragraph, text, url):
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend((properties, text_node))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text, bold=False, italic=False):
    """Add basic Markdown inline formatting to one paragraph or table cell."""
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            run.bold, run.italic = bold, italic
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.italic = italic
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(80, 80, 80)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            run.bold = bold
        elif token.startswith("["):
            label, url = re.match(r"\[([^]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        else:
            add_hyperlink(paragraph, token.rstrip(".,"), token.rstrip(".,"))
            if token[-1:] in ".,":
                paragraph.add_run(token[-1])
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        run.bold, run.italic = bold, italic


def markdown_cells(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def join_wrapped(lines):
    result = ""
    for value in lines:
        value = value.strip()
        if not result:
            result = value
        elif result.endswith("-") and value and value[0].islower():
            result += value
        else:
            result += " " + value
    return result


def configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    heading_settings = {
        "Title": (22, RGBColor(31, 78, 121)),
        "Heading 1": (16, RGBColor(31, 78, 121)),
        "Heading 2": (13, RGBColor(46, 116, 181)),
        "Heading 3": (11, RGBColor(31, 78, 121)),
    }
    for name, (size, color) in heading_settings.items():
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor(70, 70, 70)

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Consolas"
    code.font.size = Pt(8.5)
    code.paragraph_format.left_indent = Cm(0.5)
    code.paragraph_format.right_indent = Cm(0.5)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.0


def configure_document(document, title):
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    header = section.header.paragraphs[0]
    header.text = "Odos Ermou Toy Recommendation System"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Arial"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(110, 110, 110)
    add_page_number(section.footer.paragraphs[0])

    document.core_properties.title = title
    document.core_properties.subject = "Machine Learning and Content Analysis"
    document.core_properties.keywords = (
        "recommendation system, Product2Vec, Node2vec, TF-IDF, co-purchase graph"
    )
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    document.settings.element.append(update_fields)


def add_table(document, rows):
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for column, value in enumerate(rows[0]):
        cell = table.rows[0].cells[column]
        cell.text = ""
        add_inline(cell.paragraphs[0], value, bold=True)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "D9EAF7")
        set_cell_margins(cell)
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for values in rows[1:]:
        row = table.add_row()
        prevent_row_split(row)
        for column in range(len(rows[0])):
            cell = row.cells[column]
            cell.text = ""
            value = values[column] if column < len(values) else ""
            add_inline(cell.paragraphs[0], value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_image(document, image_path, caption, number):
    section = document.sections[-1]
    maximum_width = section.page_width - section.left_margin - section.right_margin
    maximum_height = Cm(17.5)
    width, height = maximum_width, None
    try:
        from PIL import Image
        with Image.open(image_path) as image:
            pixel_width, pixel_height = image.size
        aspect = pixel_width / max(pixel_height, 1)
        if maximum_width / maximum_height > aspect:
            width, height = None, maximum_height
    except (ImportError, OSError):
        pass
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(image_path), width=width, height=height)
    label = document.add_paragraph(style="Caption")
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.add_run(f"Figure {number}. {caption}")


def build_document(source_path, output_path):
    lines = source_path.read_text(encoding="utf-8").splitlines()
    title_match = next((HEADING_PATTERN.match(line) for line in lines
                        if HEADING_PATTERN.match(line)), None)
    title = title_match.group(2) if title_match else "Analytical Report"
    document = Document()
    configure_styles(document)
    configure_document(document, title)

    index = 0
    figure_number = 0
    paragraph_buffer = []
    in_toc_source = False
    title_page_open = True

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            paragraph = document.add_paragraph()
            add_inline(paragraph, join_wrapped(paragraph_buffer))
            paragraph_buffer = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        heading = HEADING_PATTERN.match(stripped)

        if in_toc_source:
            if heading and len(heading.group(1)) == 2 and heading.group(2) != "Table of Contents":
                in_toc_source = False
            else:
                index += 1
                continue

        if not stripped:
            flush_paragraph()
            index += 1
            continue

        if stripped.startswith("```"):
            flush_paragraph()
            code_lines = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            paragraph = document.add_paragraph(style="Code Block")
            paragraph.add_run("\n".join(code_lines))
            set_cell = OxmlElement("w:shd")
            set_cell.set(qn("w:fill"), "F3F4F6")
            paragraph._p.get_or_add_pPr().append(set_cell)
            index += 1
            continue

        if heading:
            flush_paragraph()
            markdown_level = len(heading.group(1))
            text = heading.group(2)
            if markdown_level == 1:
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(100)
                add_inline(paragraph, text)
            elif markdown_level == 2 and text == "Table of Contents":
                document.add_page_break()
                title_page_open = False
                paragraph = document.add_paragraph()
                paragraph.style = document.styles["Title"]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.add_run(text)
                add_toc(document.add_paragraph())
                in_toc_source = True
            else:
                if markdown_level == 2:
                    if title_page_open or len(document.paragraphs) > 1:
                        document.add_page_break()
                    title_page_open = False
                word_level = max(1, min(markdown_level - 1, 3))
                document.add_heading(text, level=word_level)
            index += 1
            continue

        if title_page_open and stripped.startswith("**") and ":**" in stripped:
            flush_paragraph()
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, stripped)
            index += 1
            continue

        image = IMAGE_PATTERN.match(stripped)
        if image:
            flush_paragraph()
            figure_number += 1
            image_path = (source_path.parent / image.group(2)).resolve()
            if not image_path.exists():
                raise FileNotFoundError(f"Report image not found: {image_path}")
            add_image(document, image_path, image.group(1), figure_number)
            index += 1
            continue

        if (stripped.startswith("|") and index + 1 < len(lines)
                and TABLE_SEPARATOR.match(lines[index + 1].strip())):
            flush_paragraph()
            table_rows = [markdown_cells(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_rows.append(markdown_cells(lines[index]))
                index += 1
            add_table(document, table_rows)
            continue

        bullet = BULLET_PATTERN.match(line)
        numbered = NUMBER_PATTERN.match(line)
        if bullet or numbered:
            flush_paragraph()
            match = bullet or numbered
            depth = min(len(match.group(1)) // 2, 2)
            base_style = "List Bullet" if bullet else "List Number"
            style = base_style if depth == 0 else f"{base_style} {depth + 1}"
            if style not in document.styles:
                style = base_style
            paragraph = document.add_paragraph(style=style)
            add_inline(paragraph, match.group(2))
            index += 1
            continue

        paragraph_buffer.append(stripped)
        index += 1

    flush_paragraph()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return {
        "output": output_path,
        "figures": figure_number,
        "tables": len(document.tables),
        "paragraphs": len(document.paragraphs),
    }


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main():
    arguments = parse_args()
    result = build_document(arguments.source.resolve(), arguments.output.resolve())
    print(f"Created {result['output']}")
    print(f"Embedded {result['figures']} figures and formatted {result['tables']} tables.")


if __name__ == "__main__":
    main()
